#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
LAB1: Генерація DOCX-звіту з JSON за шаблоном template.docx

СПРОЩЕНА ВЕРСІЯ:
- Ніякої "нормалізації" назв полів.
- Скрипт очікує, що JSON МАЄ РІВНО ТАКІ поля:

{
  "course": " ... ",
  "duration": " ... ",
  "themes": [
    {"title": "...", "durationtitle": "...", "references": "..."},
    ...
  ],
  "questions": [
    {"question": "...", "theme": "..."},
    ...
  ],
  "students": ["...", "...", ...],
  "references": ["...", "...", ...]
}

- У шаблоні template.docx мають бути плейсхолдери:
  {{course}}, {{duration}}, {{students_block}}, {{references_block}}

- Таблиці "Теми" та "Питання" заповнюються через python-docx:
  скрипт знаходить таблицю за словами в шапці:
    * теми: містить "Назва теми" і "К-сть"
    * питання: містить "Питання" і "Тема"

Встановлення:
  pip install docxtpl python-docx

Запуск (приклад):
  python generate_report.py data/report_2.json -t template.docx --outdir output

Або через run.bat (drag&drop JSON на bat).
"""

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from docxtpl import DocxTemplate
from docx import Document


# ---------------------------
# JSON
# ---------------------------

def load_json(path: Path) -> Dict[str, Any]:
    """Читає JSON і повертає dict."""
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def require_keys(data: Dict[str, Any], keys: List[str]) -> None:
    """Перевіряє наявність ключів. Якщо чогось немає — кидає помилку з поясненням."""
    missing = [k for k in keys if k not in data]
    if missing:
        raise KeyError(f"У JSON бракує обов'язкових полів: {missing}")


# ---------------------------
# DOCX helpers (tables)
# ---------------------------

def find_table_by_headers(doc: Document, headers: List[str]) -> Optional[int]:
    """
    Знаходить таблицю, якщо в 1-му рядку (шапці) є всі потрібні "headers" (підрядки).
    Повертає індекс таблиці doc.tables або None.
    """
    for i, tbl in enumerate(doc.tables):
        if not tbl.rows:
            continue
        head = " | ".join(c.text.strip() for c in tbl.rows[0].cells)
        if all(h.lower() in head.lower() for h in headers):
            return i
    return None


def clear_table_except_header(table) -> None:
    """Лишає тільки шапку (рядок 0), усі інші рядки видаляє."""
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)


def fill_themes_table(doc: Document, themes: List[Dict[str, Any]]) -> None:
    """
    Заповнює таблицю тем:
      № | Назва теми | К-сть год. | Література (№)
    JSON поля:
      themes[i].title
      themes[i].durationtitle
      themes[i].references
    """
    idx = find_table_by_headers(doc, ["Назва теми", "К-сть"])
    if idx is None:
        return  # таблицю не знайдено — нічого не робимо
    table = doc.tables[idx]

    clear_table_except_header(table)

    for i, t in enumerate(themes, start=1):
        row = table.add_row()
        cells = row.cells

        if len(cells) >= 1:
            cells[0].text = str(i)
        if len(cells) >= 2:
            cells[1].text = str(t["title"])
        if len(cells) >= 3:
            cells[2].text = str(t["durationtitle"])
        if len(cells) >= 4:
            cells[3].text = str(t["references"])


def fill_questions_table(doc: Document, questions: List[Dict[str, Any]]) -> None:
    """
    Заповнює таблицю питань:
      № | Питання | Тема
    JSON поля:
      questions[i].question
      questions[i].theme
    """
    idx = find_table_by_headers(doc, ["Питання", "Тема"])
    if idx is None:
        return
    table = doc.tables[idx]

    clear_table_except_header(table)

    for i, q in enumerate(questions, start=1):
        row = table.add_row()
        cells = row.cells

        if len(cells) >= 1:
            cells[0].text = str(i)
        if len(cells) >= 2:
            cells[1].text = str(q["question"])
        if len(cells) >= 3:
            cells[2].text = str(q["theme"])


# ---------------------------
# Rendering
# ---------------------------

def build_default_output_path(json_path: Path, out_dir: Optional[Path]) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = f"{json_path.stem}_report_{stamp}.docx"
    return (out_dir / name) if out_dir else (json_path.parent / name)


def render_docx(template_path: Path, data: Dict[str, Any], output_path: Path) -> None:
    """
    1) DocxTpl підставляє прості плейсхолдери (текст).
    2) python-docx перебудовує дві таблиці (теми + питання).
    """
    # Перетворюємо списки у багаторядкові блоки для плейсхолдерів у Word
    students = data["students"]
    refs = data["references"]
    data["students_block"] = "\n".join(f"{i}. {s}" for i, s in enumerate(students, start=1))
    data["references_block"] = "\n".join(f"{i}. {r}" for i, r in enumerate(refs, start=1))

    # 1) Підстановка плейсхолдерів
    tpl = DocxTemplate(str(template_path))
    tpl.render(data)

    tmp_path = output_path.with_suffix(".tmp.docx")
    tpl.save(str(tmp_path))

    # 2) Заповнення таблиць
    doc = Document(str(tmp_path))
    fill_themes_table(doc, data["themes"])
    fill_questions_table(doc, data["questions"])
    doc.save(str(output_path))

    # прибираємо тимчасовий файл
    try:
        tmp_path.unlink()
    except Exception:
        pass


# ---------------------------
# CLI
# ---------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="LAB1: generate DOCX report from JSON (simplified strict fields)")
    ap.add_argument("json", help="Шлях до JSON файлу (наприклад data/report_2.json)")
    ap.add_argument("-t", "--template", default="template.docx", help="Шлях до шаблону DOCX")
    ap.add_argument("-o", "--output", default=None, help="Шлях до вихідного DOCX (якщо не задано — буде auto)")
    ap.add_argument("--outdir", default=None, help="Папка виходу (якщо -o не задано)")
    args = ap.parse_args()

    root = Path(__file__).parent.resolve()

    json_path = Path(args.json).expanduser()
    if not json_path.is_absolute():
        json_path = (root / json_path).resolve()
    if not json_path.exists():
        raise SystemExit(f"JSON не знайдено: {json_path}")

    template_path = Path(args.template).expanduser()
    if not template_path.is_absolute():
        template_path = (root / template_path).resolve()
    if not template_path.exists():
        raise SystemExit(f"Шаблон не знайдено: {template_path}")

    data = load_json(json_path)

    # Строга перевірка структури
    require_keys(data, ["course", "duration", "themes", "questions", "students", "references"])

    # Строга перевірка об'єктів у themes/questions (щоб помилка була зрозуміла)
    for i, t in enumerate(data["themes"], start=1):
        if not isinstance(t, dict):
            raise TypeError(f"themes[{i}] має бути об'єктом (dict), а не {type(t)}")
        for k in ("title", "durationtitle", "references"):
            if k not in t:
                raise KeyError(f"У themes[{i}] бракує поля '{k}'")

    for i, q in enumerate(data["questions"], start=1):
        if not isinstance(q, dict):
            raise TypeError(f"questions[{i}] має бути об'єктом (dict), а не {type(q)}")
        for k in ("question", "theme"):
            if k not in q:
                raise KeyError(f"У questions[{i}] бракує поля '{k}'")

    out_dir = None
    if args.outdir:
        out_dir = Path(args.outdir).expanduser()
        if not out_dir.is_absolute():
            out_dir = (root / out_dir).resolve()
        out_dir.mkdir(parents=True, exist_ok=True)

    output_path = Path(args.output).expanduser() if args.output else build_default_output_path(json_path, out_dir)
    if not output_path.is_absolute():
        output_path = (root / output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    render_docx(template_path, data, output_path)
    print(f"OK: {output_path}")


if __name__ == "__main__":
    main()

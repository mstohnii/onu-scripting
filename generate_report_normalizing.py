#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
JSON → DOCX report generator for the provided template structure.

Pipeline:
1) docxtpl fills simple placeholders in paragraphs:
   - {{course}}, {{duration}}, {{students_block}}, {{references_block}}
2) python-docx rebuilds the two dynamic tables (Themes and Questions) from scratch:
   - find the needed table by its header row text
   - keep only the header row
   - append as many rows as needed according to JSON

Install:
  pip install docxtpl python-docx

Run:
  python generate_report.py report_2.json -t template.docx -o out.docx
"""

import argparse
import json
from pathlib import Path
from datetime import datetime
from typing import Any, Dict, List, Optional

from docxtpl import DocxTemplate
from docx import Document


# If some old/unused placeholders ever appear in template runs, we can delete them.
# (Not required for the current template, but makes the script more robust.)
PH_TEXTS = [
    "{{theme_no}}", "{{theme_title}}", "{{theme_hours}}", "{{theme_refs}}",
    "{{q_no}}", "{{q_text}}", "{{q_theme}}",
]
PH_SET = set(PH_TEXTS)


# ---------------------------
# JSON helpers
# ---------------------------

def load_json(path: Path) -> Dict[str, Any]:
    """Read JSON file and return a Python dict."""
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def normalize_data(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Normalize possible field-name variants so the rest of the code is stable.
    Your example JSON uses:
      themes[].title, themes[].durationtitle, themes[].references
      questions[].question, questions[].theme
      students[] (strings)
      references[] (strings)
    """
    data = dict(data)
    data.setdefault("course", "")
    data.setdefault("duration", "")
    data.setdefault("themes", [])
    data.setdefault("questions", [])
    data.setdefault("students", [])
    data.setdefault("references", [])

    # themes -> [{title, hours, references}]
    norm_themes: List[Dict[str, Any]] = []
    for t in data["themes"]:
        if isinstance(t, str):
            norm_themes.append({"title": t, "hours": "", "references": ""})
            continue
        t = dict(t)
        title = t.get("title") or t.get("theme") or t.get("name") or ""
        hours = t.get("hours") or t.get("duration") or t.get("durationtitle") or ""
        refs = t.get("references", "")
        if isinstance(refs, list):
            refs = ", ".join(str(x) for x in refs)
        norm_themes.append({"title": title, "hours": hours, "references": refs})
    data["themes"] = norm_themes

    # questions -> [{question, theme}]
    norm_q: List[Dict[str, Any]] = []
    for q in data["questions"]:
        if isinstance(q, str):
            norm_q.append({"question": q, "theme": ""})
            continue
        q = dict(q)
        text = q.get("question") or q.get("text") or q.get("q") or ""
        theme = q.get("theme") or q.get("themeId") or q.get("themeIndex") or ""
        norm_q.append({"question": text, "theme": theme})
    data["questions"] = norm_q

    # stringify lists (students, references)
    def stringify_list(items: List[Any]) -> List[str]:
        out: List[str] = []
        for it in items:
            if isinstance(it, str):
                out.append(it)
            elif isinstance(it, dict):
                name = it.get("name") or it.get("fullName") or it.get("student") or ""
                group = it.get("group") or ""
                out.append(f"{name} {('(' + group + ')') if group else ''}".strip())
            else:
                out.append(str(it))
        return out

    data["students"] = stringify_list(list(data["students"]))
    data["references"] = stringify_list(list(data["references"]))

    return data


# ---------------------------
# DOCX helpers
# ---------------------------

def iter_all_paragraphs(doc: Document):
    """Yield all paragraphs from body + tables + headers/footers (and their tables)."""
    for p in doc.paragraphs:
        yield p

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

    for sec in doc.sections:
        for hf in [
            sec.header, sec.footer,
            sec.first_page_header, sec.first_page_footer,
            sec.even_page_header, sec.even_page_footer,
        ]:
            if hf is None:
                continue
            for p in hf.paragraphs:
                yield p
            for tbl in hf.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p


def cleanup_stray_placeholders(doc: Document) -> None:
    """Remove placeholder-only paragraphs and wipe placeholder fragments inside runs."""
    # remove paragraphs that are exactly a placeholder
    for p in list(iter_all_paragraphs(doc)):
        if p.text.strip() in PH_SET:
            p._p.getparent().remove(p._p)

    # wipe occurrences inside runs
    for p in iter_all_paragraphs(doc):
        for r in p.runs:
            for ph in PH_TEXTS:
                if ph in r.text:
                    r.text = r.text.replace(ph, "")


def find_table_by_headers(doc: Document, headers: List[str]) -> Optional[int]:
    """
    Find a table by checking whether all required header keywords exist
    in the concatenated text of the first row.
    Returns index in doc.tables or None.
    """
    for i, tbl in enumerate(doc.tables):
        if not tbl.rows:
            continue
        head = " | ".join(c.text.strip() for c in tbl.rows[0].cells)
        if all(h.lower() in head.lower() for h in headers):
            return i
    return None


def clear_table_except_header(table) -> None:
    """Keep row 0 (header), remove the rest."""
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[1]._tr)


def fill_themes_table(doc: Document, themes: List[Dict[str, Any]]) -> None:
    """
    Fill "Themes" table. Expected columns:
      № | Назва теми | К-сть год. | Література (№)
    """
    idx = find_table_by_headers(doc, ["Назва теми", "К-сть"])
    if idx is None:
        return
    table = doc.tables[idx]
    if len(table.rows) < 1:
        return

    clear_table_except_header(table)

    for i, t in enumerate(themes, start=1):
        row = table.add_row()
        cells = row.cells
        if len(cells) >= 1:
            cells[0].text = str(i)
        if len(cells) >= 2:
            cells[1].text = str(t.get("title", ""))
        if len(cells) >= 3:
            cells[2].text = str(t.get("hours", ""))
        if len(cells) >= 4:
            cells[3].text = str(t.get("references", ""))


def fill_questions_table(doc: Document, questions: List[Dict[str, Any]]) -> None:
    """
    Fill "Questions" table. Expected columns:
      № | Питання | Тема
    """
    idx = find_table_by_headers(doc, ["Питання", "Тема"])
    if idx is None:
        return
    table = doc.tables[idx]
    if len(table.rows) < 1:
        return

    clear_table_except_header(table)

    for i, q in enumerate(questions, start=1):
        row = table.add_row()
        cells = row.cells
        if len(cells) >= 1:
            cells[0].text = str(i)
        if len(cells) >= 2:
            cells[1].text = str(q.get("question", ""))
        if len(cells) >= 3:
            cells[2].text = str(q.get("theme", ""))


# ---------------------------
# Main pipeline
# ---------------------------

def build_default_output_path(json_path: Path, out_dir: Optional[Path]) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = f"{json_path.stem}_report_{stamp}.docx"
    return (out_dir / name) if out_dir else (json_path.parent / name)


def render_docx(template_path: Path, context: Dict[str, Any], output_path: Path) -> None:
    # Build text blocks for the template placeholders
    context["students_block"] = "\n".join(
        f"{i}. {s}" for i, s in enumerate(context.get("students", []), start=1)
    )
    context["references_block"] = "\n".join(
        f"{i}. {r}" for i, r in enumerate(context.get("references", []), start=1)
    )

    # Make docxtpl tolerant to placeholders that are not in the context yet
    import jinja2

    class KeepUndefined(jinja2.Undefined):
        def __str__(self) -> str:
            return "{{" + (self._undefined_name or "") + "}}"
        __repr__ = __str__

    jinja_env = jinja2.Environment(undefined=KeepUndefined, autoescape=False)

    # 1) Render placeholders in template
    tpl = DocxTemplate(str(template_path))
    tpl.render(context, jinja_env=jinja_env)

    tmp_path = output_path.with_suffix(".tmp.docx")
    tpl.save(str(tmp_path))

    # 2) Rebuild tables via python-docx
    doc = Document(str(tmp_path))
    fill_themes_table(doc, context.get("themes", []))
    fill_questions_table(doc, context.get("questions", []))
    cleanup_stray_placeholders(doc)

    doc.save(str(output_path))

    try:
        tmp_path.unlink()
    except Exception:
        pass


def main() -> None:
    ap = argparse.ArgumentParser(description="Generate DOCX report from JSON using docxtpl + python-docx tables")
    ap.add_argument("json", help="Path to JSON file")
    ap.add_argument("-t", "--template", default="template.docx", help="Path to DOCX template")
    ap.add_argument("-o", "--output", default=None, help="Output DOCX file path")
    ap.add_argument("--outdir", default=None, help="Output directory (if -o not provided)")
    args = ap.parse_args()

    json_path = Path(args.json).expanduser().resolve()
    if not json_path.exists():
        raise SystemExit(f"JSON not found: {json_path}")

    template_path = Path(args.template).expanduser()
    if not template_path.is_absolute():
        template_path = (Path(__file__).parent.resolve() / template_path).resolve()
    if not template_path.exists():
        raise SystemExit(f"Template DOCX not found: {template_path}")

    data = normalize_data(load_json(json_path))

    out_dir = Path(args.outdir).expanduser().resolve() if args.outdir else None
    output_path = Path(args.output).expanduser().resolve() if args.output else build_default_output_path(json_path, out_dir)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    render_docx(template_path, data, output_path)
    print(f"OK: {output_path}")


if __name__ == "__main__":
    main()
